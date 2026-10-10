"""Formula cache contracts for Excel ingestion."""

from __future__ import annotations

import io
from pathlib import Path

import pandas as pd
import pytest
from apps.api.deps import settings_dep
from apps.api.main import app
from fastapi.testclient import TestClient
from mcp_servers.excel_parser import tools as excel_tools
from openpyxl import Workbook
from packages.common.config import Settings

from tests.xlsx_fixtures import inject_formula_caches

_XLSX_CT = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _formula_workbook(*, cached: bool, formula_on_other_sheet: bool = False) -> bytes:
    buffer = io.BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "data"
    sheet.append(["amount", "double"])
    sheet.append([2, "=A2*2"])
    if formula_on_other_sheet:
        other = workbook.create_sheet("other")
        other.append(["value"])
        other.append(["=1+1"])
    workbook.save(buffer)

    if not cached:
        return buffer.getvalue()

    return inject_formula_caches(
        buffer.getvalue(),
        worksheet_name="data",
        cached_values={"B2": 4},
    )


def test_uncached_formula_is_rejected_before_dataset_publish(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "uncached.xlsx"
    path.write_bytes(_formula_workbook(cached=False))

    def must_not_save(_frame: pd.DataFrame) -> str:
        raise AssertionError("workbook with uncached formulas must not publish a dataset")

    monkeypatch.setattr(excel_tools, "save_dataframe", must_not_save)
    with pytest.raises(excel_tools.FormulaCacheMissingError, match="公式.*缓存"):
        excel_tools.parse_excel({"file_ref": str(path), "sheet": "data"})


def test_cached_formula_value_is_ingested(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "cached.xlsx"
    path.write_bytes(_formula_workbook(cached=True))
    saved: list[pd.DataFrame] = []
    monkeypatch.setattr(excel_tools, "save_dataframe", lambda frame: saved.append(frame) or "ref")

    parsed = excel_tools.parse_excel({"file_ref": str(path), "sheet": "data"})

    assert parsed == {"dataset_ref": "ref", "row_count": 1, "column_count": 2}
    assert saved[0].to_dict(orient="records") == [{"amount": 2, "double": 4}]


def test_uncached_formula_on_unselected_sheet_does_not_block_ingestion(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "other-sheet.xlsx"
    path.write_bytes(_formula_workbook(cached=True, formula_on_other_sheet=True))
    saved: list[pd.DataFrame] = []
    monkeypatch.setattr(excel_tools, "save_dataframe", lambda frame: saved.append(frame) or "ref")

    excel_tools.parse_excel({"file_ref": str(path), "sheet": "data"})

    assert saved[0]["double"].tolist() == [4]


def test_uncached_formula_outside_explicit_sample_does_not_block_ingestion(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "sample.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["value"])
    sheet.append([1])
    sheet.append(["=1+1"])
    workbook.save(path)
    saved: list[pd.DataFrame] = []
    monkeypatch.setattr(excel_tools, "save_dataframe", lambda frame: saved.append(frame) or "ref")

    parsed = excel_tools.parse_excel({"file_ref": str(path), "nrows": 1})

    assert parsed["row_count"] == 1
    assert saved[0]["value"].tolist() == [1]


def test_upload_explains_uncached_formula_and_cleans_temporary_file(tmp_path: Path) -> None:
    upload_dir = tmp_path / "uploads"
    dataset_dir = tmp_path / "datasets"
    app.dependency_overrides[settings_dep] = lambda: Settings(
        upload_dir=str(upload_dir),
        dataset_dir=str(dataset_dir),
    )
    try:
        response = TestClient(app).post(
            "/upload/excel",
            files={
                "file": (
                    "uncached.xlsx",
                    _formula_workbook(cached=False),
                    _XLSX_CT,
                )
            },
        )
        assert response.status_code == 422, response.text
        assert "公式单元格缺少可用缓存值" in response.json()["detail"]
        assert "不会把这些公式静默当作空值" in response.json()["detail"]
        assert list(upload_dir.glob("*")) == []
        assert not dataset_dir.exists() or list(dataset_dir.glob("*")) == []
    finally:
        app.dependency_overrides.clear()
